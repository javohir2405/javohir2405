import os, sys, re, subprocess, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from content import load, slide_titles, FAN, YONALISH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "Malumotlar_tuzilmasi_va_algoritmlar_O‘qitish_materiallari_2026.docx".replace("‘", ""))
TNR = "Times New Roman"
TOPICS = load()
NAZ = [t for t in TOPICS if t["tur"] == "N"]
AMA = [t for t in TOPICS if t["tur"] == "A"]
assert len(NAZ) == 9 and len(AMA) == 11

# ------------------------------------------------------------------ yordamchilar
def set_font(run, size=None, bold=None, italic=None, name=TNR, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor(*color)

def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    el_pr.append(shd)

class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin = Cm(2.5), Cm(1.5)
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        st = self.d.styles["Normal"]
        st.font.name = TNR; st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing = 1.15
        for name, size, align in (("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER), ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT), ("Heading 3", 14, WD_ALIGN_PARAGRAPH.LEFT)):
            h = self.d.styles[name]
            h.font.name = TNR; h.font.size = Pt(size); h.font.bold = True; h.font.italic = (name == "Heading 3")
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
            h.element.rPr.rFonts.set(qn("w:ascii"), TNR); h.element.rPr.rFonts.set(qn("w:hAnsi"), TNR)
            h.paragraph_format.alignment = align
            h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
        self.footer()

    def footer(self):
        p = self.d.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(); set_font(r, 12)
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._element.append(f)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._element.append(it)

    # matn
    def para(self, text="", size=14, bold=False, italic=False, align=None, indent=True, space_after=0, keep=False, color=None):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "r": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "l": p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent and align is None: pf.first_line_indent = Cm(1.25)
        pf.space_after = Pt(space_after)
        if keep: pf.keep_with_next = True
        if text:
            r = p.add_run(text); set_font(r, size, bold, italic, color=color)
        return p

    def heading(self, text, level=1, page_break=False):
        h = self.d.add_heading(level=level)
        if page_break:
            h.paragraph_format.page_break_before = True
        r = h.add_run(text); set_font(r, 14, True, level == 3)
        return h

    def page_break(self):
        p = self.d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
        p.paragraph_format.space_after = Pt(0)
        return p

    def bullet(self, text, size=14):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Cm(1.25); pf.first_line_indent = Cm(-0.5); pf.space_after = Pt(2)
        r = p.add_run("–  " + text); set_font(r, size)
        return p

    def code(self, text):
        for ln in text.split("\n"):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8); pf.space_after = Pt(0); pf.line_spacing = 1.0
            shade(p._element.get_or_add_pPr(), "F2F4F7")
            r = p.add_run(ln if ln else " "); set_font(r, 10.5, name="Courier New")
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 0.5

    def table(self, rows, widths, header=True, size=12, align_center_cols=(), bold_first_col=False, shade_header="D9E2F3"):
        t = self.d.add_table(rows=0, cols=len(widths))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for ci, w in enumerate(widths):
            t.columns[ci].width = Cm(w)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci, val in enumerate(row):
                c = cells[ci]; c.width = Cm(widths[ci]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if ci in align_center_cols else WD_CELL_VERTICAL_ALIGNMENT.TOP
                lines = val if isinstance(val, list) else [val]
                first = True
                for ln in lines:
                    p = c.paragraphs[0] if first else c.add_paragraph()
                    first = False
                    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ci in align_center_cols or (header and ri == 0)) else WD_ALIGN_PARAGRAPH.LEFT
                    r = p.add_run(str(ln)); set_font(r, size, bold=(header and ri == 0) or (bold_first_col and ci == 0))
                if header and ri == 0:
                    shade(c._element.get_or_add_tcPr(), shade_header)
        if header:
            trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(6)
        return t

    def toc_line(self, text, page, level=1):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0 if level == 1 else 0.8)
        pf.right_indent = Cm(1.0)
        pf.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        pf.space_after = Pt(1); pf.line_spacing = 1.0
        r = p.add_run(f"{text}\t{page}"); set_font(r, 12 if level == 2 else 13, bold=(level == 1))
        # o'ng chekka uchun
        return p

    def save(self, path): self.d.save(path)

# ------------------------------------------------------------------ matnlar
def uz(s):  # tez yozish uchun
    return re.sub(r"(?<=[oOgG])'", "‘", s).replace("'", "’")

SOAT = 2
def tur_nomi(t): return "Nazariy" if t["tur"] == "N" else "Amaliy"

N_METOD = ["Bahs-munozara, blits-so‘rov, aqliy hujum", "Vizual ma’ruza, «Klaster» usuli, savol-javob",
           "Taqdimot, «Insert» usuli, blits-so‘rov", "Ma’ruza-suhbat, «Venn diagrammasi», aqliy hujum",
           "Muammoli ma’ruza, «Tushunchalar tahlili», mini-munozara"]
A_METOD = ["Amaliy ish, juftlikda dasturlash, o‘zaro tekshiruv", "Keys-stadi, kichik guruhlarda ishlash, laboratoriya ishi",
           "Amaliy topshiriq, «Muammoli vaziyat» usuli, taqdimot", "Loyiha usuli, kichik guruhlarda ishlash, hisobot himoyasi"]

# ------------------------------------------------------------------ bo‘limlar
def cover(D):
    c = "c"
    D.para("O‘ZBEKISTON RESPUBLIKASI", 14, True, align=c)
    D.para("OLIY TA’LIM, FAN VA INNOVATSIYALAR VAZIRLIGI", 14, True, align=c, space_after=10)
    D.para("______________________________________", 14, True, align=c)
    D.para("(ta’lim muassasasi nomi)", 11, align=c, space_after=14)
    p = D.para("“TASDIQLAYMAN”", 14, True, align="r")
    D.para("Ta’lim muassasasi rahbari", 14, align="r")
    D.para("________ ____________________", 14, align="r")
    D.para("“___”__________ 2026 yil", 14, align="r", space_after=40)
    D.para("MA’LUMOTLAR TUZILMASI VA ALGORITMLAR", 18, True, align=c)
    D.para("FANIDAN", 16, True, align=c, space_after=6)
    D.para("O‘QITISH MATERIALLARI", 22, True, align=c, space_after=30)
    D.para(f"Ta’lim yo‘nalishi:  “{YONALISH}”", 14, align="l", indent=False, space_after=4)
    D.para("Fan bo‘yicha jami soat:  40 soat (nazariy – 18 soat, amaliy – 22 soat)", 14, align="l", indent=False, space_after=24)
    D.para("O‘qitish materiallari ta’lim muassasasi pedagogik kengashining 2026 yildagi ___-sonli yig‘ilishida o‘quv jarayoniga qo‘llashga tavsiya etildi.", 14, indent=False, space_after=8)
    D.para("O‘qitish materiallari to‘plami “_______________________” kafedrasining 2026 yildagi ___-sonli yig‘ilishida muhokama etildi va ma’qullandi.", 14, indent=False, space_after=14)
    D.para("Kafedra mudiri: _________________________      ____________", 14, align="l", indent=False)
    D.para("(F.I.SH.)                                                                                        (imzo)", 10.5, align="l", indent=False, space_after=8)
    D.para("O‘qituvchi:       _________________________      ____________", 14, align="l", indent=False)
    D.para("(F.I.SH.)                                                                                        (imzo)", 10.5, align="l", indent=False, space_after=36)
    D.para("____________ – 2026", 14, True, align=c)

def author(D):
    D.heading("TO‘PLAMNING MUALLIFI TO‘G‘RISIDA MA’LUMOT", 1, page_break=True)
    D.para("", space_after=4)
    rows = [["Fotosurat (3×4)\n\n\n\n", ""],]
    t = D.d.add_table(rows=1, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = t.rows[0].cells
    c0.width = Cm(4.2); c1.width = Cm(12.3)
    t.autofit = False
    t.columns[0].width = Cm(4.2); t.columns[1].width = Cm(12.3)
    c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c0.paragraphs[0].add_run("\n\n3×4\nfotosurat\n\n"); set_font(r, 12, italic=True, color=(120, 120, 120))
    info = ["O‘qituvchining F.I.SH.: ______________________________", "Ma’lumoti: ______________________________",
            "Mutaxassisligi: ______________________________", "Ilmiy darajasi va unvoni: ______________________",
            "Pedagogik ish staji: ______ yil", "Ish joyi: ______________________________",
            "Elektron pochta manzili: ______________________", "Telefon raqami: ______________________"]
    first = True
    for ln in info:
        p = c1.paragraphs[0] if first else c1.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(ln); set_font(r, 13)
    D.para("", space_after=10)
    D.para("O‘QITISH MATERIALLARI TO‘PLAMI TARKIBI:", 14, True, align="l", indent=False, space_after=6)
    D.table([["№", "To‘plam tarkibi", "Shakli"],
             ["1", "O‘qitish materiallari to‘plami titul varag‘i", "Matn"],
             ["2", "To‘plamning muallifi to‘g‘risida ma’lumot", "Matn"],
             ["3", "Namunaviy o‘quv dasturi", "Matn, jadval"],
             ["4", "Ishchi o‘quv dasturi (18 soat nazariy, 22 soat amaliy)", "Matn, jadval"],
             ["5", "Taqvim-mavzular rejasi (40 soat)", "Jadval"],
             ["6", "Har bir mavzu bo‘yicha texnologik xaritalar", "Jadval"],
             ["7", "Mavzular to‘plami (20 ta mavzu matni)", "Yozma + elektron (Word)"],
             ["8", "Vizual-didaktik resurslar (har bir mavzuga 10 varaqli slayd)", "PowerPoint (.pptx)"],
             ["9", "Baholash mezoni", "Matn, jadval"],
             ["10", "Foydalanilgan adabiyotlar ro‘yxati", "Ro‘yxat"]],
            [1.2, 11.3, 4.0], align_center_cols=(0, 2), size=13)

NAZ_RESULTS = [
    "Ma’lumotlar tuzilmalari va algoritmlarning asosiy tushunchalarini, ularning tasnifini va murakkabligini baholash usullarini;",
    "massiv, bog‘langan ro‘yxat, stek, navbat, daraxt, graf va xesh jadvalning tuzilishi, amallari va qo‘llanilish sohalarini;",
    "qidirish va saralash algoritmlarining g‘oyasi, afzalliklari va kamchiliklarini;",
    "rekursiya va «bo‘lib tashla va hukmronlik qil» tamoyilining mohiyatini;",
    "kutubxona va axborot tizimlarida ma’lumotlar tuzilmalarining o‘rnini bilishi kerak.",
]
KONIKMA = [
    "algoritmni so‘z, blok-sxema, psevdokod va dasturlash tilida ifodalash;",
    "masalaga mos ma’lumotlar tuzilmasi va algoritmni asoslab tanlash;",
    "algoritmning vaqt va xotira murakkabligini baholash;",
    "tuzilma va algoritmlarni Python dasturlash tilida amalga oshirish hamda sinash;",
    "kutubxona ma’lumotlarini (katalog, buyurtmalar, tarmoqlar) qayta ishlash ko‘nikmalariga ega bo‘lishi kerak.",
]

def namunaviy(D):
    D.heading("NAMUNAVIY O‘QUV DASTURI", 1, page_break=True)
    D.para(f"Fan nomi: “{FAN}”", 14, True, align="l", indent=False)
    D.para(f"Ta’lim yo‘nalishi: {YONALISH}", 14, align="l", indent=False)
    D.para("Fanga ajratilgan jami soat: 40 (nazariy – 18, amaliy – 22)", 14, align="l", indent=False, space_after=8)
    D.heading("1. Fanning maqsadi va vazifalari", 2)
    D.para("Fanni o‘qitishdan maqsad – talabalarda ma’lumotlarni tashkil etish va qayta ishlashning asosiy tuzilmalari hamda algoritmlari bo‘yicha nazariy bilim va amaliy ko‘nikmalarni shakllantirish, ularni kutubxona va axborot tizimlari faoliyatiga tatbiq eta olish qobiliyatini rivojlantirishdir.")
    D.para("Fanning vazifalari: algoritm va ma’lumotlar tuzilmasi tushunchalarini o‘rgatish; algoritmlarning samaradorligini baholash usullarini o‘zlashtirish; asosiy tuzilmalar (massiv, ro‘yxat, stek, navbat, daraxt, graf, xesh jadval) bilan ishlash ko‘nikmasini hosil qilish; qidirish va saralash algoritmlarini o‘rganish; masalaga mos tuzilma va algoritmni tanlashga o‘rgatish; algoritmik fikrlashni rivojlantirish.")
    D.heading("2. Talabalarning bilim, ko‘nikma va malakalariga qo‘yiladigan talablar", 2)
    D.para("Fanni o‘zlashtirish natijasida talaba:", indent=False)
    D.para("bilishi kerak:", bold=True, align="l", indent=False)
    for x in NAZ_RESULTS: D.bullet(x)
    D.para("ko‘nikmalarga ega bo‘lishi kerak:", bold=True, align="l", indent=False)
    for x in KONIKMA: D.bullet(x)
    D.para("malakaga ega bo‘lishi kerak:", bold=True, align="l", indent=False)
    D.bullet("kichik hajmdagi axborot tizimini (masalan, kutubxona katalogi) mustaqil loyihalash va dasturlash;")
    D.bullet("o‘z yechimining to‘g‘riligi va samaradorligini asoslab, himoya qilish malakalariga ega bo‘lishi kerak.")
    D.heading("3. Fanning mazmuni", 2)
    D.para("Nazariy mashg‘ulotlar mazmuni (18 soat):", bold=True, align="l", indent=False, keep=True)
    rows = [["№", "Mavzu nomi va qisqacha mazmuni", "Soat"]]
    for i, t in enumerate(NAZ, 1):
        rows.append([str(i), [t["title"] + ".", "Mazmuni: " + "; ".join(h[0].lower() + h[1:] for h, _ in t["plans"]) + "."], "2"])
    rows.append(["", "Jami:", "18"])
    D.table(rows, [1.2, 14.0, 1.3], align_center_cols=(0, 2), size=12)
    D.para("Amaliy mashg‘ulotlar mazmuni (22 soat):", bold=True, align="l", indent=False, keep=True)
    rows = [["№", "Mavzu nomi va qisqacha mazmuni", "Soat"]]
    for i, t in enumerate(AMA, 1):
        rows.append([str(i), [t["title"] + ".", "Mazmuni: " + "; ".join(h[0].lower() + h[1:] for h, _ in t["plans"]) + "."], "2"])
    rows.append(["", "Jami:", "22"])
    D.table(rows, [1.2, 14.0, 1.3], align_center_cols=(0, 2), size=12)
    D.heading("4. Fanni o‘qitish shakllari va usullari", 2)
    D.para("Fan bo‘yicha nazariy mashg‘ulotlar vizual ma’ruza, savol-javob, aqliy hujum, blits-so‘rov, «Klaster», «Insert», «Venn diagrammasi» kabi interfaol usullar yordamida o‘tkaziladi. Amaliy mashg‘ulotlar kompyuter xonasida, har bir talaba uchun alohida kompyuter bilan, juftlikda dasturlash, kichik guruhlarda ishlash, keys-stadi va loyiha usullari asosida tashkil etiladi. Dasturlash tili sifatida Python tanlangan, chunki u o‘rganish uchun sodda va algoritmlarni ixcham ifodalash imkonini beradi.")
    D.heading("5. Baholash tizimi", 2)
    D.para("Talabalar bilimi joriy, oraliq va yakuniy nazorat shaklida 100 ballik tizimda baholanadi. Baholash mezonlari to‘plamning 9-bo‘limida batafsil keltirilgan.")
    D.para("", space_after=10)
    D.para("Namunaviy dasturni tuzuvchi:  _______________________   ____________", 13, align="l", indent=False)
    D.para("Kelishildi:  ______________________________   ____________", 13, align="l", indent=False)

def ishchi(D):
    D.heading("ISHCHI O‘QUV DASTURI", 1, page_break=True)
    D.para(f"Fan: “{FAN}”.  Jami: 40 soat, shundan nazariy – 18 soat, amaliy – 22 soat.", 14, align="l", indent=False, space_after=6)
    rows = [["№", "Mavzu nomi", "Jami soat", "Nazariy", "Amaliy"]]
    for t in TOPICS:
        rows.append([str(t["no"]), t["title"], "2", "2" if t["tur"] == "N" else "–", "2" if t["tur"] == "A" else "–"])
    rows.append(["", "JAMI:", "40", "18", "22"])
    D.table(rows, [1.1, 9.3, 1.9, 2.1, 2.1], align_center_cols=(0, 2, 3, 4), size=12.5)
    D.para("Mavzularning o‘quv rejalari (har bir mavzu 4 ta reja savolidan iborat):", 14, True, align="l", indent=False, keep=True)
    rows = [["№", "Mavzu va reja savollari"]]
    for t in TOPICS:
        rows.append([str(t["no"]), [f"{t['title']} ({tur_nomi(t).lower()})"] + [f"{i}. {h}" for i, (h, _) in enumerate(t["plans"], 1)]])
    D.table(rows, [1.1, 15.4], align_center_cols=(0,), size=12)

def taqvim(D):
    D.heading("TAQVIM-MAVZULAR REJASI (40 soat)", 1, page_break=True)
    D.para(f"Fan: “{FAN}”.  Nazariy – 18 soat, amaliy – 22 soat.  Haftada 4 soat (2 ta juft mashg‘ulot), jami 10 hafta.", 13, align="l", indent=False, space_after=6)
    rows = [["№", "Mashg‘ulot mavzusi", "Turi", "Soat", "Hafta", "O‘tkazish sanasi"]]
    for t in TOPICS:
        rows.append([str(t["no"]), t["title"], tur_nomi(t), "2", str((t["no"] + 1) // 2), "___.___.2026"])
    rows.append(["", "JAMI:", "", "40", "", ""])
    D.table(rows, [1.0, 8.2, 2.0, 1.2, 1.5, 2.6], align_center_cols=(0, 2, 3, 4, 5), size=12)
    rows = [["Soatlar taqsimoti", "Nazariy", "Amaliy", "Jami"],
            ["Mashg‘ulotlar soni", "9", "11", "20"],
            ["Soat", "18", "22", "40"],
            ["Foiz", "45 %", "55 %", "100 %"]]
    D.table(rows, [6.5, 3.3, 3.3, 3.4], align_center_cols=(1, 2, 3), size=12.5, bold_first_col=True)
    D.para("Eslatma: sanalar o‘quv yilining dars jadvaliga muvofiq o‘qituvchi tomonidan to‘ldiriladi.", 12, italic=True, align="l", indent=False)

def tex_xarita(D):
    D.heading("O‘QUV MASHG‘ULOTLARINING TEXNOLOGIK XARITALARI", 1, page_break=True)
    D.para("Quyida fanning har bir mavzusi uchun o‘quv mashg‘ulotining ta’lim texnologiyasi modeli va texnologik xaritasi keltirilgan. Har bir mashg‘ulot 2 soat (80 daqiqa) davom etadi.", 14)
    for t in TOPICS:
        N = t["tur"] == "N"
        D.heading(f"{t['no']}-mavzu: {t['title']} — texnologik xarita", 2, page_break=(t["no"] > 1))
        D.para("O‘quv mashg‘ulotining ta’lim texnologiyasi modeli", 13, True, align="c", indent=False, space_after=4, keep=True)
        plans = [f"{i}. {h}" for i, (h, _) in enumerate(t["plans"], 1)]
        if N:
            shakl = "Nazariy: yangi bilimlarni egallash bo‘yicha o‘quv mashg‘uloti"
            maqsad = f"“{t['title']}” mavzusi bo‘yicha ma’lumot berish, asosiy tushunchalarni shakllantirish va mavzuni mustahkamlash."
            ped = [f"{i}. {h} haqida tushuncha beradi;" for i, (h, _) in enumerate(t["plans"], 1)]
            nat = [f"{i}. {h} haqida bilib oladilar;" for i, (h, _) in enumerate(t["plans"], 1)]
            metod = N_METOD[(t["no"] // 2) % len(N_METOD)]
            vosita = "Tarqatma materiallar, slaydlar orqali taqdimot, proyektor, kompyuter"
            sharoit = "Proyektor orqali taqdimot o‘tkazishga ixtisoslashtirilgan auditoriya"
            fb = "Og‘zaki savol-javob, blits-so‘rov, nazorat savollari"
            shakl2 = "Jamoa bilan va kichik guruhlarda ishlash"
        else:
            shakl = "Amaliy: bilim, ko‘nikma va malakalarni shakllantirish bo‘yicha o‘quv mashg‘uloti"
            maqsad = f"“{t['title']}” mavzusi bo‘yicha nazariy bilimlarni amalda qo‘llash, dasturlash ko‘nikma va malakalarini shakllantirish."
            ped = [f"{i}. {h} bo‘yicha topshiriq beradi va bajarilishini yo‘naltiradi;" for i, (h, _) in enumerate(t["plans"], 1)]
            nat = [f"{i}. {h} bo‘yicha topshiriqni bajaradilar;" for i, (h, _) in enumerate(t["plans"], 1)]
            metod = A_METOD[(t["no"] // 2) % len(A_METOD)]
            vosita = "Kompyuterlar, Python dasturlash muhiti (Thonny/IDLE), tarqatma topshiriqlar, proyektor"
            sharoit = "Kompyuter xonasi (har bir o‘quvchi uchun alohida kompyuter)"
            fb = "Topshiriqlarni tekshirish, o‘zaro tekshiruv, hisobotni himoya qilish"
            shakl2 = "Juftlikda va kichik guruhlarda ishlash"
        rows = [
            ["Vaqt: 80 daqiqa (2 soat)", "O‘quvchilar soni: 10 nafar"],
            ["O‘quv mashg‘ulotining shakli va turi", shakl],
            ["O‘quv mashg‘uloti rejasi", plans],
            ["O‘quv mashg‘uloti maqsadi", maqsad],
            ["Pedagogik vazifalar", ped],
            ["O‘quv faoliyati natijalari", nat],
            ["O‘qitish metodlari", metod],
            ["O‘quv faoliyatini tashkil etish shakli", shakl2],
            ["O‘qitish vositalari", vosita],
            ["O‘qitish sharoiti", sharoit],
            ["Qayta aloqaning usul va vositalari", fb],
        ]
        D.table(rows, [5.0, 11.5], header=False, size=11.5, bold_first_col=True)
        D.para("O‘quv mashg‘ulotining texnologik xaritasi", 13, True, align="c", indent=False, space_after=4, keep=True)
        R = [["Ish bosqichlari", "Faoliyat mazmuni: o‘qituvchi", "Faoliyat mazmuni: o‘quvchi"]]
        R.append(["1-bosqich. O‘quv mashg‘ulotiga kirish (5 daqiqa)",
                  ["Tashkiliy qism:", "1. O‘quvchilarning davomatini tekshiradi.", "2. Mashg‘ulotga tayyorgarlikni nazorat qiladi.",
                   "3. Mavzuning nomi, maqsadi, rejasi va kutilayotgan natijalar haqida ma’lumot beradi."],
                  ["Mashg‘ulotga tayyorgarlik ko‘radilar.", "Tinglaydilar, mavzu nomini va rejani daftarga yozadilar."]])
        if N:
            teacher = ["Tayanch bilimlarni faollashtirish:",
                       "2.1. Oldingi mavzu bo‘yicha “Blits-so‘rov” usulida savollar beradi (4 daqiqa).",
                       "Yangi o‘quv materiali bayoni:"]
            teacher += [f"2.{i+1}. {i}-reja: “{h}” bo‘yicha ma’ruza qiladi, slaydlarni namoyish etadi, qisqa muhokama asosida umumlashtiradi (13 daqiqa)." for i, (h, _) in enumerate(t["plans"], 1)]
            teacher += ["Yangi o‘quv materialini mustahkamlash:", "2.6. “Aqliy hujum” usulida savollar beradi, g‘oyalarni doskaga yozib, umumlashtiradi (9 daqiqa)."]
            pupil = ["Savollarga javob beradilar.", "Berilgan ma’lumotlarni daftarlarida qayd qilib boradilar.", "Har bir reja yakunidagi muhokamada ishtirok etadilar.",
                     "Savollar beradilar, o‘z fikrlarini bayon etadilar.", "Aqliy hujumda g‘oyalar aytadilar."]
            last_t = ["Mavzu bo‘yicha yakuniy xulosa chiqaradi.", "Faol o‘quvchilarni rag‘batlantiradi va baholaydi.",
                      "Uyga vazifa beradi: nazorat savollariga yozma javob tayyorlash va keyingi amaliy mashg‘ulotga tayyorlanish."]
        else:
            teacher = ["Tayanch bilimlarni faollashtirish:",
                       "2.1. Oldingi nazariy mavzu bo‘yicha savol-javob o‘tkazadi, topshiriq shartlarini va baholash mezonlarini tushuntiradi (5 daqiqa).",
                       "Amaliy topshiriqlarni bajarish:"]
            teacher += [f"2.{i+1}. {i}-bosqich: “{h}” bo‘yicha topshiriqni beradi, bajarilishini kuzatadi, yo‘naltiradi (14 daqiqa)." for i, (h, _) in enumerate(t["plans"], 1)]
            teacher += ["Natijalarni mustahkamlash:", "2.6. Guruhlar yechimlarini o‘zaro tekshirishni tashkil etadi, xatolarni tahlil qiladi (4 daqiqa)."]
            pupil = ["Savollarga javob beradilar.", "Topshiriq shartini daftarga yozadilar, kompyuterda dasturlash muhitini tayyorlaydilar.",
                     "Topshiriqlarni juftlikda yoki guruhda bajaradilar, kodni sinab ko‘radilar.", "Natijalarni tahlil qiladilar, savollar beradilar.",
                     "Hamkasbining yechimini tekshiradilar."]
            last_t = ["Natijalarni taqdimot qildiradi, guruhlar ishini baholaydi.", "Umumiy xulosa chiqaradi, xatolarni tahlil qiladi.",
                      "Uyga vazifa beradi: topshiriqni yakunlash va hisobot tayyorlash."]
        R.append(["2-bosqich. Asosiy bosqich (65 daqiqa)", teacher, pupil])
        R.append(["3-bosqich. Yakuniy bosqich (10 daqiqa)", last_t,
                  ["Tinglaydilar, o‘z fikrlarini bayon etadilar.", "Savollar beradilar.", "Uyga vazifani yozib oladilar."]])
        D.table(R, [3.4, 8.5, 4.6], header=True, size=11.5, bold_first_col=True)

def mavzu_matnlari(D, markers):
    D.heading("MAVZULAR TO‘PLAMI (MA’RUZA MATNLARI)", 1, page_break=True)
    D.para("Quyida fanning 20 ta mavzusi bo‘yicha o‘quv materiallari matnlari keltirilgan. Har bir mavzu 4 ta reja savolidan iborat bo‘lib, dasturlash misollari Python tilida berilgan.", 14)
    for t in TOPICS:
        h = D.heading(f"{t['no']}-MAVZU: {t['title']}", 2, page_break=True)
        D.para(f"({tur_nomi(t)} mashg‘ulot, 2 soat)", 12, italic=True, align="c", indent=False, space_after=4)
        D.para("Mavzu rejasi:", 14, True, align="l", indent=False, keep=True)
        for i, (hh, _) in enumerate(t["plans"], 1):
            D.para(f"{i}. {hh}", 14, align="l", indent=False)
        D.para("", space_after=4)
        for i, (hh, blocks) in enumerate(t["plans"], 1):
            D.heading(f"{i}. {hh}", 3)
            for kind, x in blocks:
                if kind == "p": D.para(x)
                elif kind == "li": D.bullet(x)
                elif kind == "code": D.code(x)
        D.heading("Nazorat savollari", 3)
        for _, q in t["questions"]:
            D.bullet(q)

SLIDE_DIR = "/tmp/claude-0/-home-user-javohir2405/ae8cccd8-7417-56ec-96d8-2472d0742e4b/scratchpad/slides"

def slaydlar(D):
    D.heading("VIZUAL-DIDAKTIK RESURSLAR (SLAYDLAR)", 1, page_break=True)
    D.para("Har bir mavzu uchun mavzu matnidan olingan 10 varaqli taqdimot (slayd) tayyorlangan. Slaydlar quyida mavzular bo‘yicha keltirilgan; ularning PowerPoint (.pptx) variantlari ham alohida fayllarda mavjud. Har bir taqdimot tarkibi: 1 – titul; 2 – mavzu rejasi; 3–9 – reja savollari bo‘yicha mazmun va dasturlash namunalari; 10 – xulosa va nazorat savollari.", 14)
    for t in TOPICS:
        D.heading(f"{t['no']}-mavzu slaydlari: {t['title']}", 2, page_break=True)
        tb = D.d.add_table(rows=5, cols=2); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i in range(10):
            c = tb.rows[i // 2].cells[i % 2]
            pr = c.paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pr.paragraph_format.space_after = Pt(4); pr.paragraph_format.space_before = Pt(4)
            pr.add_run().add_picture(os.path.join(SLIDE_DIR, f"s{t['no']:02d}-{i+1:02d}.jpg"), width=Cm(7.9))
        for row in tb.rows:
            trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)

TEST = [
    ("Algoritmning chekliligi xossasi nimani anglatadi?", ["A) Algoritm chekli sondagi qadamdan so‘ng to‘xtashi kerak", "B) Algoritm faqat chekli son bilan ishlaydi", "C) Algoritm cheklangan xotira talab qiladi", "D) Algoritm bitta natija beradi"], "A"),
    ("n elementli tartiblangan massivda binar qidirishning murakkabligi qanday?", ["A) O(1)", "B) O(log n)", "C) O(n)", "D) O(n²)"], "B"),
    ("Qaysi saralash algoritmi barqaror va har qanday holatda O(n log n) vaqt oladi?", ["A) Pufakchali saralash", "B) Tez saralash", "C) Birlashtirish orqali saralash", "D) Tanlash orqali saralash"], "C"),
    ("Bir bog‘lamli ro‘yxatning boshiga element qo‘shishning murakkabligi qanday?", ["A) O(1)", "B) O(log n)", "C) O(n)", "D) O(n²)"], "A"),
    ("Stek qaysi tamoyil bo‘yicha ishlaydi?", ["A) FIFO", "B) LIFO", "C) Tasodifiy tanlash", "D) Ustuvorlik bo‘yicha"], "B"),
    ("Rekursiv funksiyaning to‘xtashini ta’minlovchi qism nima deb ataladi?", ["A) Rekursiv holat", "B) Asosiy holat", "C) Chaqiruv freymi", "D) Memoizatsiya"], "B"),
    ("Ikkilik qidiruv daraxtini simmetrik tartibda aylanib chiqsak, qiymatlar qanday tartibda chiqadi?", ["A) Kamayish tartibida", "B) O‘sish tartibida", "C) Tasodifiy tartibda", "D) Daraja bo‘yicha"], "B"),
    ("Graf bo‘ylab kenglikka birinchi qidirish (BFS) qaysi tuzilmadan foydalanadi?", ["A) Stek", "B) Navbat", "C) Xesh jadval", "D) Massiv"], "B"),
    ("Xesh jadvalda ikki xil kalit bir xil chelakka tushishi nima deb ataladi?", ["A) Qayta xeshlash", "B) Yuklanish", "C) To‘qnashuv", "D) Probalash"], "C"),
    ("Qaysi tuzilma kutubxona katalogida ISBN bo‘yicha o‘rtacha O(1) vaqtda qidirish uchun eng mos?", ["A) Bog‘langan ro‘yxat", "B) Xesh jadval", "C) Stek", "D) Navbat"], "B"),
]

def baholash(D):
    D.heading("BAHOLASH MEZONI", 1, page_break=True)
    D.para("Talabalarning fan bo‘yicha bilim, ko‘nikma va malakalari 100 ballik reyting tizimida joriy, oraliq va yakuniy nazorat turlari orqali baholanadi.", 14)
    D.heading("1. Reyting ballarining taqsimlanishi", 2)
    D.table([["Nazorat turi", "Mazmuni", "Maksimal ball"],
             ["Joriy nazorat", ["11 ta amaliy mashg‘ulot topshirig‘i (har biri 3 ball) – 33 ball", "Nazariy mashg‘ulotlardagi faollik, blits-so‘rovlar va davomat – 7 ball"], "40"],
             ["Oraliq nazorat", "Yozma ish yoki test (nazariy mavzular va dasturlash masalalari)", "20"],
             ["Yakuniy nazorat", ["Yakuniy loyiha («Kutubxona axborot tizimi») va uni himoya qilish – 25 ball", "Yakuniy test – 15 ball"], "40"],
             ["JAMI", "", "100"]],
            [3.4, 10.6, 2.5], align_center_cols=(0, 2), size=12.5, bold_first_col=True)
    D.heading("2. Baholash shkalasi", 2)
    D.table([["Ball (foiz)", "Baho", "Bilim darajasi"],
             ["86 – 100", "5 (a’lo)", "Mavzuni chuqur o‘zlashtirgan, mustaqil fikrlay oladi, murakkab masalalarni yecha oladi va o‘z yechimini asoslaydi"],
             ["71 – 85", "4 (yaxshi)", "Mavzuni to‘g‘ri o‘zlashtirgan, standart masalalarni mustaqil yechadi, ba’zi kamchiliklarga yo‘l qo‘yadi"],
             ["55 – 70", "3 (qoniqarli)", "Asosiy tushunchalarni biladi, masalalarni o‘qituvchi yordamida yechadi"],
             ["0 – 54", "2 (qoniqarsiz)", "Mavzuni o‘zlashtirmagan, masalalarni yecha olmaydi"]],
            [3.0, 3.2, 10.3], align_center_cols=(0, 1), size=12.5)
    D.heading("3. Amaliy topshiriqni baholash mezoni (1 ta mashg‘ulot uchun – 3 ball)", 2)
    D.table([["Mezon", "Ball", "Izoh"],
             ["Yechimning to‘g‘riligi", "1,5", "Dastur barcha sinov misollarida to‘g‘ri natija beradi, chegaraviy holatlar hisobga olingan"],
             ["Kod sifati", "0,5", "Nomlar ma’noli, izohlar mavjud, kod tushunarli va tartibli yozilgan"],
             ["Sinovlar va murakkablik tahlili", "0,5", "Sinov misollari yetarli, algoritmning vaqt va xotira murakkabligi asoslangan"],
             ["Hisobot va himoya", "0,5", "Hisobot to‘liq, talaba yechimini tushuntira oladi va savollarga javob beradi"]],
            [5.0, 1.5, 10.0], align_center_cols=(1,), size=12.5)
    D.heading("4. Yakuniy loyihani baholash mezoni (25 ball)", 2)
    D.table([["Mezon", "Ball"],
             ["Ma’lumotlar tuzilmalarini to‘g‘ri va asosli tanlash", "7"],
             ["Dasturning to‘liq va to‘g‘ri ishlashi", "6"],
             ["Kod sifati, hujjatlashtirish va avtomatik sinovlar", "4"],
             ["Murakkablik tahlili va yaxshilash yo‘llari", "4"],
             ["Taqdimot va savollarga javoblar", "4"]],
            [13.0, 3.5], align_center_cols=(1,), size=12.5)
    D.heading("5. Namunaviy test savollari", 2)
    for i, (q, opts, ans) in enumerate(TEST, 1):
        D.para(f"{i}. {q}", 13, True, align="l", indent=False, keep=True)
        for o in opts: D.para(o, 13, align="l", indent=False)
    D.para("Javoblar kaliti: " + ", ".join(f"{i}-{a}" for i, (_, _, a) in enumerate(TEST, 1)) + ".", 13, bold=True, align="l", indent=False)
    D.para("Har bir mavzuning oxirida keltirilgan nazorat savollari og‘zaki so‘rov va oraliq nazorat savollarini tuzishda ishlatiladi.", 13, italic=True, align="l", indent=False)

ADABIYOT = {
    "Normativ-huquqiy hujjatlar": [
        "O‘zbekiston Respublikasining «Ta’lim to‘g‘risida»gi Qonuni. 2020 yil 23 sentabr, O‘RQ-637-son.",
        "O‘zbekiston Respublikasining «Axborotlashtirish to‘g‘risida»gi Qonuni. 2003 yil 11 dekabr.",
        "O‘zbekiston Respublikasining «Axborot-kutubxona faoliyati to‘g‘risida»gi Qonuni. 2011 yil 13 aprel.",
        "«Raqamli O‘zbekiston – 2030» strategiyasini tasdiqlash va uni samarali amalga oshirish chora-tadbirlari to‘g‘risida. O‘zbekiston Respublikasi Prezidentining 2020 yil 5 oktabrdagi PF-6079-son Farmoni.",
    ],
    "Asosiy adabiyotlar": [
        "Cormen T.H., Leiserson C.E., Rivest R.L., Stein C. Introduction to Algorithms. 4th ed. – Cambridge, MA: MIT Press, 2022. – 1312 p.",
        "Sedgewick R., Wayne K. Algorithms. 4th ed. – Upper Saddle River, NJ: Addison-Wesley, 2011. – 969 p.",
        "Goodrich M.T., Tamassia R., Goldwasser M.H. Data Structures and Algorithms in Python. – Hoboken, NJ: Wiley, 2013. – 748 p.",
        "Skiena S.S. The Algorithm Design Manual. 3rd ed. – Cham: Springer, 2020. – 793 p.",
        "Кормен Т., Лейзерсон Ч., Ривест Р., Штайн К. Алгоритмы: построение и анализ. 3-е изд. – М.: Вильямс, 2013. – 1328 с.",
        "Вирт Н. Алгоритмы и структуры данных. – М.: ДМК Пресс, 2010.",
    ],
    "Qo‘shimcha adabiyotlar": [
        "Knuth D.E. The Art of Computer Programming. Vol. 1: Fundamental Algorithms. 3rd ed. – Reading, MA: Addison-Wesley, 1997.",
        "Knuth D.E. The Art of Computer Programming. Vol. 3: Sorting and Searching. 2nd ed. – Reading, MA: Addison-Wesley, 1998.",
        "Wirth N. Algorithms + Data Structures = Programs. – Englewood Cliffs, NJ: Prentice-Hall, 1976. – 366 p.",
        "Aho A.V., Hopcroft J.E., Ullman J.D. Data Structures and Algorithms. – Reading, MA: Addison-Wesley, 1983. – 427 p.",
        "Bhargava A. Grokking Algorithms. – Shelter Island, NY: Manning, 2016. – 256 p.",
        "[O‘quv rejada ko‘rsatilgan o‘zbek tilidagi darslik va o‘quv qo‘llanmalar bu yerga qo‘shiladi]",
    ],
    "Internet resurslari": [
        "Python 3 hujjatlari: https://docs.python.org/3/",
        "VisuAlgo – algoritmlar vizualizatsiyasi: https://visualgo.net",
        "Algorithms for Competitive Programming: https://cp-algorithms.com",
        "Khan Academy – Algorithms: https://www.khanacademy.org/computing/computer-science/algorithms",
        "O‘zbekiston Respublikasi qonun hujjatlari ma’lumotlari milliy bazasi: https://lex.uz",
        "Ziyo NET axborot-ta’lim portali: https://ziyonet.uz",
    ],
}

def adabiyot(D):
    D.heading("FOYDALANILGAN ADABIYOTLAR RO‘YXATI", 1, page_break=True)
    n = 0
    for sec, items in ADABIYOT.items():
        D.para(sec, 14, True, align="l", indent=False, space_after=4, keep=True)
        for it in items:
            n += 1
            p = D.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.first_line_indent = Cm(-1.0); p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f"{n}.  {it}"); set_font(r, 14)
        D.para("", space_after=4)

# ------------------------------------------------------------------ yig‘ish
SECTIONS = [
    ("TO‘PLAMNING MUALLIFI TO‘G‘RISIDA MA’LUMOT", 1),
    ("NAMUNAVIY O‘QUV DASTURI", 1),
    ("ISHCHI O‘QUV DASTURI", 1),
    ("TAQVIM-MAVZULAR REJASI (40 soat)", 1),
    ("O‘QUV MASHG‘ULOTLARINING TEXNOLOGIK XARITALARI", 1),
    ("MAVZULAR TO‘PLAMI (MA’RUZA MATNLARI)", 1),
] + [(f"{t['no']}-MAVZU: {t['title']}", 2) for t in TOPICS] + [
    ("VIZUAL-DIDAKTIK RESURSLAR (SLAYDLAR)", 1),
    ("BAHOLASH MEZONI", 1),
    ("FOYDALANILGAN ADABIYOTLAR RO‘YXATI", 1),
]

def build(pages):
    D = Doc()
    cover(D)
    D.heading("MUNDARIJA", 1, page_break=True)
    D.para("", space_after=4)
    for (txt, lvl) in SECTIONS:
        D.toc_line(txt, pages.get(txt, "0"), lvl)
    author(D); namunaviy(D); ishchi(D); taqvim(D); tex_xarita(D); mavzu_matnlari(D, None); slaydlar(D); baholash(D); adabiyot(D)
    D.save(OUT)
    return OUT

def find_pages(pdf):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout).group(1))
    texts = []
    for i in range(1, n + 1):
        t = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), pdf, "-"], capture_output=True, text=True).stdout
        texts.append(re.sub(r"\s+", " ", t))
    pages = {}
    start = 2  # titul + mundarijadan keyin
    for (txt, lvl) in SECTIONS:
        key = re.sub(r"\s+", " ", txt)[:45]
        for i in range(start, n):
            if key in texts[i]:
                pages[txt] = str(i + 1)
                break
    return pages, n

if __name__ == "__main__":
    tmp = "/tmp/claude-0/-home-user-javohir2405/ae8cccd8-7417-56ec-96d8-2472d0742e4b/scratchpad/docbuild"
    os.makedirs(tmp, exist_ok=True)
    pages = {}
    for it in range(2):
        out = build(pages)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, out], capture_output=True, timeout=600)
        pdf = os.path.join(tmp, os.path.basename(out).replace(".docx", ".pdf"))
        pages, n = find_pages(pdf)
        print("pass", it, "pages:", n, "found:", len(pages), "of", len(SECTIONS))
    missing = [s for s, _ in SECTIONS if s not in pages]
    print("missing:", missing)
    print(out)
