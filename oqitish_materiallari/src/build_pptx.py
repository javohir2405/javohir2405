import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from content import load, plan_slides, FAN, YONALISH

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xE0, 0x8E, 0x2B)
INK = RGBColor(0x22, 0x2B, 0x38)
MUTE = RGBColor(0x5B, 0x67, 0x78)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
CODEBG = RGBColor(0xF1, 0xF4, 0xF8)
W, H = Inches(13.333), Inches(7.5)
FONT = "Calibri"

def add_text(slide, x, y, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return tb

def base(prs, title, no, total=10):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.35))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.35), W, Inches(0.07))
    acc.fill.solid(); acc.fill.fore_color.rgb = ACCENT; acc.line.fill.background()
    size = 30 if len(title) < 55 else (26 if len(title) < 85 else 22)
    add_text(s, Inches(0.5), Inches(0.1), Inches(12.3), Inches(1.15), title, size, PAPER, True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(11.6), Inches(7.0), Inches(1.5), Inches(0.35), f"{no} / {total}", 12, MUTE, align=PP_ALIGN.RIGHT)
    return s

def bullets(slide, items, size=24):
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(1.75), Inches(11.9), Inches(5.1))
    tf = tb.text_frame; tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        r0 = p.add_run(); r0.text = "■  "; r0.font.size = Pt(size - 8); r0.font.color.rgb = ACCENT; r0.font.name = FONT
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = FONT

def build(topic, out, prs=None):
    own = prs is None
    if own:
        prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    slides = plan_slides(topic)
    tur = "Nazariy mashg‘ulot" if topic["tur"] == "N" else "Amaliy mashg‘ulot"
    for n, sd in enumerate(slides, 1):
        k = sd["kind"]
        if k == "title":
            s = prs.slides.add_slide(prs.slide_layouts[6])
            bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
            bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
            ac = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.55), Inches(2.2), Inches(0.1))
            ac.fill.solid(); ac.fill.fore_color.rgb = ACCENT; ac.line.fill.background()
            add_text(s, Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.5), f"{FAN} fani", 20, RGBColor(0xC9,0xD6,0xE8))
            add_text(s, Inches(0.8), Inches(1.4), Inches(11.7), Inches(2.1), f"{topic['no']}-MAVZU. {topic['title']}", 38 if len(topic['title']) < 60 else 32, PAPER, True, anchor=MSO_ANCHOR.BOTTOM)
            add_text(s, Inches(0.8), Inches(3.9), Inches(11), Inches(0.5), tur + " · 2 soat (80 daqiqa)", 22, RGBColor(0xF3,0xC1,0x7A))
            add_text(s, Inches(0.8), Inches(6.3), Inches(11.5), Inches(0.8), f"Yo‘nalish: {YONALISH}", 16, RGBColor(0xC9,0xD6,0xE8))
        elif k == "plan":
            s = base(prs, "Mavzu rejasi", n)
            bullets(s, [f"{i+1}. {h}" for i, (h, _) in enumerate(topic["plans"])], 26)
        elif k == "bullets":
            s = base(prs, sd["title"], n); bullets(s, sd["bullets"], 24 if sum(len(b) for b in sd["bullets"]) < 520 else 22)
        elif k == "code":
            s = base(prs, sd["title"], n)
            lines = sd["code"].split("\n")[:16]
            box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.75), Inches(11.9), Inches(5.0))
            box.fill.solid(); box.fill.fore_color.rgb = CODEBG; box.line.color.rgb = RGBColor(0xD0,0xD7,0xE2)
            tf = box.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.15)
            size = 18 if len(lines) <= 10 else (16 if len(lines) <= 13 else 14)
            for i, ln in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT
                r = p.add_run(); r.text = ln if ln else " "
                r.font.size = Pt(size); r.font.name = "Courier New"; r.font.color.rgb = INK
        elif k == "questions":
            s = base(prs, "Xulosa va nazorat savollari", n)
            qs = [x for _, x in topic["questions"]][:5]
            bullets(s, qs, 22)
    if own:
        prs.save(out)
    return slides

if __name__ == "__main__":
    outdir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "slaydlar")
    os.makedirs(outdir, exist_ok=True)
    for t in load():
        fn = os.path.join(outdir, f"{t['no']:02d}-mavzu_slaydlar.pptx")
        build(t, fn)
        print("ok", fn)

    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    for t in load():
        build(t, None, prs)
    out = os.path.join(os.path.dirname(outdir), "Malumotlar_tuzilmasi_va_algoritmlar_Slaydlar_20_mavzu.pptx")
    prs.save(out); print(len(prs.slides), out)
