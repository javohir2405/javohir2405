import sys
SP = sys.argv[1]
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
TNR = 'Times New Roman'

def font(run, size=14, bold=False, italic=False):
    run.font.name = TNR; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), TNR)

def para(doc_or_cell, text='', size=14, bold=False, align=None, indent=None, space_after=6, italic=False):
    p = doc_or_cell.add_paragraph()
    if text: font(p.add_run(text), size, bold, italic)
    pf = p.paragraph_format; pf.space_after = Pt(space_after); pf.space_before = Pt(0)
    if align is not None: p.alignment = align
    if indent is not None: pf.first_line_indent = Cm(indent)
    return p

def borders(table):
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '6'); x.set(qn('w:color'), '000000'); b.append(x)
    table._tbl.tblPr.append(b)

def cell_text(cell, text, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    font(p.add_run(text), size, bold)
    cell.vertical_alignment = 1

def new_doc():
    doc = docx.Document()
    sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(1.5); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    st = doc.styles['Normal']; st.font.name = TNR; st.font.size = Pt(14)
    return doc

# ---------------- 10) Baholash mezoni
doc = new_doc()
para(doc, 'Baholash mezoni', 14, True, WD_ALIGN_PARAGRAPH.LEFT, space_after=10)
para(doc, 'O‘quvchilar baholash mezoni va baholash ko‘rsatkichlari bilan tanishtirib o‘tiladi:', 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.25, 14)
tb = doc.add_table(rows=5, cols=4); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(tb); tb.autofit = False
widths = [Cm(7.2), Cm(2.8), Cm(2.4), Cm(3.6)]
for ci, w_ in enumerate(widths): tb.columns[ci].width = w_
head = tb.rows[0].cells[0].merge(tb.rows[0].cells[3])
cell_text(head, 'O‘quvchilar faoliyatini baholash mexanizmi', 14, True)
rows = [
 ('Berilgan savol yoki topshiriqni to‘liq bajarish yoki bayon etish', '86-100%', '5 baho', 'A’lo'),
 ('Berilgan savol yoki topshiriqni imloviy xatolar, to‘xtalishlar bilan bajarish, bayon etish', '71-85%', '4 baho', 'Yaxshi'),
 ('Topshiriqni bajarishda texnik va dasturiy xatoliklarga yo‘l qo‘yish', '56-70%', '3 baho', 'Qoniqarli'),
 ('Topshiriqni bajarishda qo‘pol xatoliklarga yo‘l qo‘yish yoki javob bera olmaslik', '55% dan past', '2 baho', 'Qoniqarsiz'),
]
for ri, r in enumerate(rows, 1):
    for ci, v in enumerate(r):
        c = tb.rows[ri].cells[ci]; c.width = widths[ci]
        cell_text(c, v, 14, False, WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER)
doc.save(f'{SP}/out/part_baholash.docx')

# ---------------- 11) Foydalanilgan adabiyotlar
doc = new_doc()
para(doc, 'Foydalanilgan adabiyotlar ro‘yxati', 14, True, WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
def section(title, items, start):
    para(doc, title, 14, True, WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
    for k, it in enumerate(items, start):
        p = para(doc, '%d. %s' % (k, it), 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6)
        p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.9)
    return start + len(items)
n = 1
n = section('Asosiy adabiyotlar', [
 'Mo‘minov B.B. Dasturlash 1. Darslik. – T.: “Nihol print”, 2021. – 280 b.',
 'Mo‘minov B.B. Dasturlash 2. Darslik. – T.: “Nihol print”, 2021. – 604 b.',
 'Nazirov Sh.A., Qobulov R.V., Babajanov M.R., Raxmanov Q.S. C va C++ tili. – T.: “Voris-nashriyot” MChJ, 2013. – 488 b.',
 'Kutubxonashunoslik va bibliografiya (50320203) mutaxassisligi bo‘yicha “Dasturlash” moduli (5PM0187) namunaviy o‘quv dasturi. – T.: Professional ta’limni rivojlantirish instituti, 2025.',
], n)
n = section('Qo‘shimcha adabiyotlar', [
 'Stroustrup B. The C++ Programming Language. 4th ed. – Upper Saddle River: Addison-Wesley, 2013.',
 'Stroustrup B. Programming: Principles and Practice Using C++. 2nd ed. – Upper Saddle River: Addison-Wesley, 2014.',
 'Josuttis N.M. The C++ Standard Library: A Tutorial and Reference. 2nd ed. – Upper Saddle River: Addison-Wesley, 2012.',
 'Troelsen A., Japikse P. Pro C# 9 with .NET 5: Foundational Principles and Practices in Programming. 10th ed. – New York: Apress, 2021.',
], n)
n = section('Internet resurslari', [
 'Microsoft Learn. Windows Forms hujjatlari (Windows Forms documentation): https://learn.microsoft.com/dotnet/desktop/winforms/',
 'Microsoft Learn. Visual Studio’da C++ hujjatlari (C++ in Visual Studio documentation): https://learn.microsoft.com/cpp/',
 'cppreference.com. C++ Reference (C++ standart kutubxonasi): https://en.cppreference.com/',
 'ZiyoNET axborot-ta’lim portali: https://ziyonet.uz/',
], n)
doc.save(f'{SP}/out/part_adabiyot.docx')
print('ok')
